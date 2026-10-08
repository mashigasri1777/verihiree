"""
VERIRESUME - Sample Data Generator
Generates realistic PDF and DOCX resume files for testing all academic scenarios.
"""

import sys
from pathlib import Path
import fitz  # PyMuPDF
import docx  # python-docx

# Ensure project root is in sys.path
BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import config

SAMPLES_OUT_DIR = config.SAMPLE_DIR
STATIC_SAMPLES_DIR = BASE_DIR / "static" / "samples"
SAMPLES_OUT_DIR.mkdir(parents=True, exist_ok=True)
STATIC_SAMPLES_DIR.mkdir(parents=True, exist_ok=True)


def create_pdf_resume(file_name: str, content_blocks: list) -> Path:
    """Generates a styled PDF resume using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4

    # Top header background banner
    rect = fitz.Rect(0, 0, 595, 75)
    page.draw_rect(rect, color=(0.04, 0.07, 0.17), fill=(0.04, 0.07, 0.17))

    y = 35
    # Name
    page.insert_text((40, y), content_blocks[0]["text"], fontsize=18, fontname="helv", color=(0.0, 0.94, 1.0))
    y += 22
    # Subtitle / Contact
    page.insert_text((40, y), content_blocks[1]["text"], fontsize=9, fontname="helv", color=(0.9, 0.9, 0.9))

    y = 105
    for block in content_blocks[2:]:
        b_type = block.get("type", "p")
        text = block.get("text", "")

        if b_type == "h2":
            y += 12
            page.draw_line((40, y + 14), (555, y + 14), color=(0.8, 0.8, 0.8), width=0.5)
            page.insert_text((40, y + 10), text, fontsize=12, fontname="helv", color=(0.0, 0.47, 0.71))
            y += 26
        elif b_type == "bullet":
            page.insert_text((55, y), "•  " + text, fontsize=9.5, fontname="helv", color=(0.15, 0.15, 0.15))
            y += 16
        else:
            page.insert_text((40, y), text, fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.2))
            y += 15

    out_path = SAMPLES_OUT_DIR / file_name
    doc.save(str(out_path))
    doc.close()

    # Also save to static/samples
    static_path = STATIC_SAMPLES_DIR / file_name
    static_path.write_bytes(out_path.read_bytes())
    return out_path


def create_docx_resume(file_name: str, candidate_name: str, email: str, certs: list) -> Path:
    """Generates a DOCX resume using python-docx."""
    doc = docx.Document()

    # Header
    title = doc.add_heading(candidate_name, level=0)
    p_contact = doc.add_paragraph(f"{email} | (555) 019-4829 | San Francisco, CA")

    # Skills
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph("Python, FastAPI, SQL, Cloud Computing, Machine Learning, Git, Unit Testing")

    # Education
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science, State University, 2020 – 2024")

    # Certifications
    doc.add_heading("CERTIFICATIONS & CREDENTIALS", level=1)
    for c in certs:
        doc.add_paragraph(
            f"{c['name']} | Certificate ID: {c['id']} | Issued by: {c['org']} | Date: {c['date']}",
            style="List Bullet"
        )

    out_path = SAMPLES_OUT_DIR / file_name
    doc.save(str(out_path))

    static_path = STATIC_SAMPLES_DIR / file_name
    static_path.write_bytes(out_path.read_bytes())
    return out_path


def generate_all_samples():
    """Generates all scenario sample files."""
    print("Generating Academic Test Resumes...")

    # 1. Alice Johnson (Scenario 1 - 100% Match Valid)
    create_pdf_resume("sample_alice_johnson_valid.pdf", [
        {"type": "name", "text": "ALICE JOHNSON"},
        {"type": "contact", "text": "alice.johnson@example.com | (555) 019-2834 | San Francisco, CA"},
        {"type": "h2", "text": "PROFESSIONAL SUMMARY"},
        {"type": "p", "text": "Dedicated Software Engineer with 3+ years experience in Python backend development."},
        {"type": "h2", "text": "TECHNICAL SKILLS"},
        {"type": "p", "text": "Python, FastAPI, Docker, PostgreSQL, REST APIs, Microservices, SymPy, Multiprocessing"},
        {"type": "h2", "text": "EDUCATION"},
        {"type": "p", "text": "Bachelor of Science in Computer Science, Demo University, 2020 – 2024"},
        {"type": "h2", "text": "CERTIFICATIONS"},
        {"type": "bullet", "text": "Python Programming | Certificate ID: CERT1001 | Issued by Demo Institute | Date: 2024-05-15"},
    ])

    # 2. Bob Smith (Scenario 2 - Name Mismatch Suspicious)
    create_pdf_resume("sample_bob_smith_mismatch.pdf", [
        {"type": "name", "text": "JOHN DOE"},
        {"type": "contact", "text": "john.doe@example.com | (555) 012-3456 | New York, NY"},
        {"type": "h2", "text": "PROFESSIONAL SUMMARY"},
        {"type": "p", "text": "Full Stack Web Developer specializing in modern frontend and backend architectures."},
        {"type": "h2", "text": "TECHNICAL SKILLS"},
        {"type": "p", "text": "JavaScript, React, Node.js, Python, CSS3, HTML5, MongoDB"},
        {"type": "h2", "text": "EDUCATION"},
        {"type": "p", "text": "Bachelor of Technology in Information Technology, City College, 2019 – 2023"},
        {"type": "h2", "text": "CERTIFICATIONS"},
        {"type": "bullet", "text": "Web Development | Certificate ID: CERT1002 | Issued by Demo Institute | Date: 2024-06-20"},
    ])

    # 3. Multi-Certificate Resume (Scenario 4 - Multiprocessing Demo with 4 certs)
    create_pdf_resume("sample_multi_cert_resume.pdf", [
        {"type": "name", "text": "ALICE JOHNSON"},
        {"type": "contact", "text": "alice.j@example.com | (555) 777-8888 | Seattle, WA"},
        {"type": "h2", "text": "PROFESSIONAL SUMMARY"},
        {"type": "p", "text": "Senior Cloud and AI Solutions Architect with multiple verified industry credentials."},
        {"type": "h2", "text": "TECHNICAL SKILLS"},
        {"type": "p", "text": "Cloud Architecture, AWS, Google Cloud, Deep Learning, Python, Kubernetes"},
        {"type": "h2", "text": "EDUCATION"},
        {"type": "p", "text": "Master of Science in Computer Science, Stanford University"},
        {"type": "h2", "text": "CERTIFICATIONS & CREDENTIALS"},
        {"type": "bullet", "text": "Python Programming | Certificate ID: CERT1001 | Issued by Demo Institute | Date: 2024-05-15"},
        {"type": "bullet", "text": "AWS Certified Solutions Architect - Associate | Certificate ID: AWS-SAA-8842 | Issued by Amazon Web Services | Date: 2024-01-18"},
        {"type": "bullet", "text": "Deep Learning Specialization | Certificate ID: DEEP-AI-991 | Issued by DeepLearning.AI | Date: 2024-03-22"},
        {"type": "bullet", "text": "CS50: Introduction to Computer Science | Certificate ID: CS50-HARVARD-2023 | Issued by Harvard University | Date: 2023-12-15"},
    ])

    # 4. Revoked Certificate Resume (Scenario 3 - Revocation Check)
    create_pdf_resume("sample_revoked_cert_resume.pdf", [
        {"type": "name", "text": "FAKE CANDIDATE"},
        {"type": "contact", "text": "fake.candidate@example.com | (555) 000-9999"},
        {"type": "h2", "text": "SUMMARY"},
        {"type": "p", "text": "Demonstration profile containing a flagged/revoked certificate ID."},
        {"type": "h2", "text": "CERTIFICATIONS"},
        {"type": "bullet", "text": "Python Programming | Certificate ID: CERT9999 | Issued by Demo Institute | Date: 2022-01-01"},
    ])

    # 5. DOCX Sample Resume
    create_docx_resume(
        "sample_resume_charlie.docx",
        "David Miller",
        "david.miller@example.com",
        [
            {"name": "Google Professional Data Engineer", "id": "GCP-PDE-3011", "org": "Google Cloud", "date": "2023-09-05"}
        ]
    )

    print("All sample resumes created successfully in sample_data/ and static/samples/")


if __name__ == "__main__":
    generate_all_samples()
